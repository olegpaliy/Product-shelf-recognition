"""JSON and Excel report writers."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict

from openpyxl import Workbook


def write_json(path: Path, payload: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def write_excel(path: Path, payload: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    summary = wb.active
    summary.title = "summary"
    qa = payload.get("qa", {})
    plan = payload.get("planogram", {})
    counts = payload.get("brand_counts", {})
    share = payload.get("brand_share", {})

    rows = [
        ("image", payload.get("image")),
        ("result_id", payload.get("result_id")),
        ("valid_photo", qa.get("valid")),
        ("qa_reason", qa.get("reason")),
        ("blur_score", qa.get("blur_score")),
        ("brightness", qa.get("brightness")),
        ("detections", payload.get("detection_count")),
        ("presence_pct", plan.get("presence_pct")),
        ("row_match_pct", plan.get("row_match_pct")),
        ("coarse_compliance_pct", plan.get("coarse_compliance_pct")),
        ("missing", ", ".join(plan.get("missing") or [])),
        ("competitor", ", ".join(plan.get("competitor") or [])),
    ]
    summary.append(["metric", "value"])
    for k, v in rows:
        summary.append([k, v])

    dist = wb.create_sheet("numeric_distribution")
    dist.append(["brand", "count", "share_pct"])
    for brand, count in counts.items():
        dist.append([brand, count, share.get(brand, 0)])

    det = wb.create_sheet("detections")
    det.append(["idx", "brand", "confidence", "shelf_row", "x1", "y1", "x2", "y2", "det_conf"])
    for i, item in enumerate(payload.get("detections", []), start=1):
        det.append(
            [
                i,
                item.get("brand"),
                item.get("brand_confidence"),
                item.get("shelf_row"),
                item.get("x1"),
                item.get("y1"),
                item.get("x2"),
                item.get("y2"),
                item.get("confidence"),
            ]
        )

    shelves = wb.create_sheet("shelves")
    shelves.append(["shelf", "brands_json", "total"])
    for row in payload.get("shelf_summary", []):
        shelves.append([row.get("shelf"), json.dumps(row.get("brands", {})), row.get("total")])

    path_out = Path(path)
    wb.save(path_out)
